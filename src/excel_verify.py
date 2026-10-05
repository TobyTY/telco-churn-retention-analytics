"""Recalculate the workbook in Excel (when available) and check formula results against facts.json.

Without Excel (CI, Linux) recalculation is skipped and the check confirms the KPI cells hold formulas.
"""
from __future__ import annotations

import sys

from openpyxl import load_workbook

from src import facts
from src.config import EXCEL_PATH


def expected(fx: dict) -> dict:
    k, g = fx["analysis"]["kpi"], fx["model"]["campaign"]["groups"]["model"]
    return {"B5": (k["customers"], 0), "B6": (k["churned"], 0), "B7": (k["churn_rate"], 1e-6), "B8": (k["mrr_lost"], 0.05),
            "B9": (k["annual_revenue_lost"], 0.5), "B10": (g["customers"], 0), "B11": (g["net_value"], 5.0),
            "B12": (g["roi"], 1e-4), "B13": (g["breakeven_save_rate"], 1e-4),
            "B14": (fx["model"]["risk_bands_summary"][0]["expected_mrr_at_risk"], 0.05)}


def recalc_with_excel() -> bool:
    try:
        import win32com.client  # type: ignore
        xl = win32com.client.DispatchEx("Excel.Application")
    except Exception:  # noqa: BLE001 - no Excel on this machine
        return False
    xl.Visible, xl.DisplayAlerts = False, False
    try:
        wb = xl.Workbooks.Open(str(EXCEL_PATH.resolve()))
        xl.CalculateFull()
        wb.Save()
        wb.Close(SaveChanges=False)
    finally:
        xl.Quit()
    return True


def main() -> int:
    recalculated = recalc_with_excel()
    fx = facts.read()
    formulas = load_workbook(EXCEL_PATH)["Summary"]
    values = load_workbook(EXCEL_PATH, data_only=True)["Summary"]
    bad = []
    for cell, (exp, tol) in expected(fx).items():
        if not str(formulas[cell].value).startswith("="):
            bad.append(f"{cell} is not a formula")
            continue
        got = values[cell].value
        if recalculated and (got is None or abs(float(got) - float(exp)) > max(tol, 1e-9)):
            bad.append(f"{cell}: Excel {got} vs facts {exp}")
    facts.update("excel_check", {"recalculated_by_excel": recalculated, "cells_checked": len(expected(fx)), "mismatches": len(bad)})
    if bad:
        print("excel_verify: FAIL\n  " + "\n  ".join(bad))
        return 1
    print(f"excel_verify: OK - {len(expected(fx))} KPI cells " + ("recalculated in Excel and matched facts.json" if recalculated else "are formulas (no Excel)"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
