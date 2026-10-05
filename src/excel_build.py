"""Build excel/retention_campaign_roi.xlsx: a formula-driven retention campaign ROI calculator.

Sheets: Summary, ROI Calculator (inputs with data validation, group selector, sensitivity table and
chart), Groups, Segments (churn rate and Wilson 95% CI as Excel formulas), Risk bands, Lift.
src/excel_verify.py recalculates the workbook and checks it against facts.json.
"""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from src import facts
from src.config import EXCEL_PATH, REPORTS, SQL_OUT

NAVY = "1F2A44"
HEAD, HEAD_FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY)
INPUT_FILL, BASE_FILL = PatternFill("solid", fgColor="FFF4CC"), PatternFill("solid", fgColor="E8F1FC")
USD, USD2, PCT, PCT2, INT = '$#,##0', '$#,##0.00', '0.0%', '0.00%', '#,##0'
GROUP_LABELS = {"model": "Model: churn probability >= threshold", "rule_m2m_fiber_first_year": "Rule: month-to-month + fiber + first year",
                "all_month_to_month": "All month-to-month customers", "everyone": "Everyone"}


def header(ws, row, cols, start=1):
    for j, c in enumerate(cols, start):
        cell = ws.cell(row, j, c)
        cell.font, cell.fill = HEAD, HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def title(ws, text, sub):
    ws["A1"], ws["A2"] = text, sub
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"].font = Font(italic=True, size=9, color="52514E")


def widths(ws, w):
    for k, v in w.items():
        ws.column_dimensions[k].width = v


def build() -> None:
    fx = facts.read()
    m, a = fx["model"], fx["analysis"]
    camp = m["campaign_assumptions"]
    wb = Workbook()

    # ---------------------------------------------------------------- Groups
    wg = wb.active
    wg.title = "Groups"
    title(wg, "Candidate target groups", "Values from the out-of-fold model scores (src/model.py). Churn rate is a formula.")
    header(wg, 4, ["Group", "Customers", "Churners", "Churn rate", "Churner MRR ($)", "Group MRR ($)"])
    for i, (k, lab) in enumerate(GROUP_LABELS.items(), 5):
        g = m["campaign"]["groups"][k]
        wg.cell(i, 1, lab)
        wg.cell(i, 2, g["customers"]).number_format = INT
        wg.cell(i, 3, g["churners"]).number_format = INT
        wg.cell(i, 4, f"=C{i}/B{i}").number_format = PCT
        wg.cell(i, 5, round(g["churner_mrr"], 2)).number_format = USD2
        wg.cell(i, 6, round(g["group_mrr"], 2)).number_format = USD2
    wg.cell(10, 1, f"Model threshold used: {m['threshold']['chosen']:.2f} (chosen by maximum net value on training folds)").font = Font(italic=True, size=9)
    widths(wg, {"A": 44, "B": 11, "C": 10, "D": 11, "E": 15, "F": 15})

    # ---------------------------------------------------------------- ROI calculator
    wc = wb.create_sheet("ROI Calculator", 0)
    title(wc, "Retention campaign ROI calculator", "Yellow cells are inputs. Pick a target group in B5; every result is a formula.")
    inputs = [
        ("Target group", GROUP_LABELS["model"], None),
        ("Discount on monthly bill", camp["offer_discount_pct"] / 100, PCT),
        ("Discount length (months)", camp["offer_months"], INT),
        ("Contact cost per targeted customer ($)", camp["contact_cost"], USD2),
        ("Save rate: share of would-be churners retained", camp["save_rate"], PCT),
        ("Revenue horizon for a saved customer (months)", camp["value_horizon_months"], INT),
    ]
    wc.cell(4, 1, "Inputs").font = Font(bold=True, color=NAVY)
    for r, (lab, val, fmt) in enumerate(inputs, 5):
        wc.cell(r, 1, lab)
        c = wc.cell(r, 2, val)
        c.fill = INPUT_FILL
        if fmt:
            c.number_format = fmt
    dv_group = DataValidation(type="list", formula1="=Groups!$A$5:$A$8", allow_blank=False)
    dv_pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", showErrorMessage=True, error="Enter 0%-100%")
    dv_int = DataValidation(type="whole", operator="between", formula1="1", formula2="60", showErrorMessage=True, error="Enter 1-60 months")
    dv_cost = DataValidation(type="decimal", operator="between", formula1="0", formula2="500", showErrorMessage=True, error="Enter $0-500")
    for dv, ref in [(dv_group, "B5"), (dv_pct, "B6"), (dv_int, "B7"), (dv_cost, "B8"), (dv_pct, "B9"), (dv_int, "B10")]:
        wc.add_data_validation(dv)
        dv.add(ref)
    wc.cell(12, 1, "Selected group").font = Font(bold=True, color=NAVY)
    look = [("Customers targeted", 2, INT), ("Would-be churners in group", 3, INT), ("Churner MRR ($)", 5, USD2), ("Group MRR ($)", 6, USD2)]
    for r, (lab, col, fmt) in enumerate(look, 13):
        wc.cell(r, 1, lab)
        c = wc.cell(r, 2, f"=INDEX(Groups!{chr(64 + col)}$5:{chr(64 + col)}$8,MATCH($B$5,Groups!$A$5:$A$8,0))")
        c.number_format, c.fill = fmt, BASE_FILL
    wc.cell(18, 1, "Results").font = Font(bold=True, color=NAVY)
    results = [
        ("Customers saved", "=B14*B9", "0.0"),
        ("Revenue saved over horizon ($)", "=B9*B15*B10", USD),
        ("Campaign cost: contact + discount ($)", "=B8*B13+B6*B7*B16", USD),
        ("Net value ($)", "=B20-B21", USD),
        ("ROI (net / cost)", "=IF(B21=0,0,B22/B21)", PCT),
        ("Break-even save rate", "=IF(B15=0,\"n/a\",B21/(B15*B10))", PCT),
        ("Cost per customer saved ($)", "=IF(B19=0,\"n/a\",B21/B19)", USD2),
    ]
    for r, (lab, fml, fmt) in enumerate(results, 19):
        wc.cell(r, 1, lab)
        c = wc.cell(r, 2, fml)
        c.number_format, c.font = fmt, Font(bold=True)
    wc.conditional_formatting.add("B22", CellIsRule(operator="lessThan", formula=["0"], font=Font(bold=True, color="D03B3B")))
    wc.conditional_formatting.add("B22", CellIsRule(operator="greaterThanOrEqual", formula=["0"], font=Font(bold=True, color="006300")))

    # sensitivity: net value by save rate (rows) x discount (columns) for the selected group
    wc.cell(28, 1, "Sensitivity: net value ($) by save rate and discount, selected group").font = Font(bold=True, color=NAVY)
    discounts = [0.10, 0.15, 0.20, 0.25, 0.30]
    rates = [0.10, 0.20, 0.30, 0.40, 0.50]
    header(wc, 29, ["Save rate \\ discount"] + [f"{d:.0%}" for d in discounts])
    for j, d in enumerate(discounts, 2):
        wc.cell(30, j, d).number_format = PCT
    wc.cell(30, 1, "discount ->").font = Font(italic=True, size=9)
    for i, s in enumerate(rates, 31):
        wc.cell(i, 1, s).number_format = PCT
        for j in range(2, 2 + len(discounts)):
            col = chr(64 + j)
            wc.cell(i, j, f"=$A{i}*$B$15*$B$10-($B$8*$B$13+{col}$30*$B$7*$B$16)").number_format = USD
    wc.conditional_formatting.add("B31:F35", ColorScaleRule(start_type="min", start_color="D03B3B", mid_type="num", mid_value=0,
                                                            mid_color="FFFFFF", end_type="max", end_color="2A78D6"))
    ch = LineChart()
    ch.title, ch.y_axis.title, ch.x_axis.title, ch.height, ch.width = "Net value by save rate (20% discount)", "Net value ($)", "Save rate", 7, 13
    ch.add_data(Reference(wc, min_col=4, min_row=30, max_row=35), titles_from_data=True)
    ch.set_categories(Reference(wc, min_col=1, min_row=31, max_row=35))
    wc.add_chart(ch, "D4")
    widths(wc, {"A": 46, "B": 40, "C": 12, "D": 12, "E": 12, "F": 12})

    # ---------------------------------------------------------------- Segments with Wilson CI formulas
    ws = wb.create_sheet("Segments")
    title(ws, "Churn by segment with Wilson 95% confidence intervals", "Values: customers, churned, MRR lost. Rate, CI and shares are formulas; z in B3.")
    ws["A3"], ws["B3"] = "z (95%)", "=NORM.S.INV(0.975)"
    seg = pd.read_csv(REPORTS / "segment_churn.csv")
    header(ws, 4, ["Dimension", "Segment", "Customers", "Churned", "MRR lost ($)", "Churn rate", "CI low", "CI high", "Share of MRR lost"])
    first = 5
    for i, r in enumerate(seg.itertuples(), first):
        ws.cell(i, 1, r.dimension)
        ws.cell(i, 2, str(r.segment))
        ws.cell(i, 3, int(r.customers)).number_format = INT
        ws.cell(i, 4, int(r.churned)).number_format = INT
        ws.cell(i, 5, round(r.mrr_lost, 2)).number_format = USD
        ws.cell(i, 6, f"=D{i}/C{i}").number_format = PCT
        ws.cell(i, 7, f"=(F{i}+$B$3^2/(2*C{i})-$B$3*SQRT(F{i}*(1-F{i})/C{i}+$B$3^2/(4*C{i}^2)))/(1+$B$3^2/C{i})").number_format = PCT
        ws.cell(i, 8, f"=(F{i}+$B$3^2/(2*C{i})+$B$3*SQRT(F{i}*(1-F{i})/C{i}+$B$3^2/(4*C{i}^2)))/(1+$B$3^2/C{i})").number_format = PCT
        ws.cell(i, 9, f"=E{i}/SUMIF($A${first}:$A${first + len(seg) - 1},A{i},$E${first}:$E${first + len(seg) - 1})").number_format = PCT
    last = first + len(seg) - 1
    ws.conditional_formatting.add(f"F{first}:F{last}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="D03B3B"))
    ws.auto_filter.ref = f"A4:I{last}"
    ws.freeze_panes = "A5"
    widths(ws, {"A": 18, "B": 26, "C": 11, "D": 10, "E": 13, "F": 11, "G": 10, "H": 10, "I": 15})
    contract_rows = [i for i, r in enumerate(seg.itertuples(), first) if r.dimension == "contract"]

    # ---------------------------------------------------------------- Risk bands
    wr = wb.create_sheet("Risk bands")
    title(wr, "Model risk bands (out-of-fold scores)", "High >= 50% predicted churn, Medium 25-50%, Low < 25%. Shares are formulas.")
    rb = pd.read_csv(SQL_OUT / "risk_bands.csv")
    header(wr, 4, ["Risk band", "Customers", "Actual churn rate", "Avg predicted probability", "MRR ($)", "Expected MRR at risk ($)",
                   "Share of customers", "Share of expected risk"])
    for i, r in enumerate(rb.itertuples(), 5):
        wr.cell(i, 1, r.risk_band)
        wr.cell(i, 2, int(r.customers)).number_format = INT
        wr.cell(i, 3, round(r.actual_churn_rate, 6)).number_format = PCT
        wr.cell(i, 4, round(r.avg_probability, 6)).number_format = PCT
        wr.cell(i, 5, round(r.mrr, 2)).number_format = USD
        wr.cell(i, 6, round(r.expected_mrr_at_risk, 2)).number_format = USD
        wr.cell(i, 7, f"=B{i}/SUM($B$5:$B$7)").number_format = PCT
        wr.cell(i, 8, f"=F{i}/SUM($F$5:$F$7)").number_format = PCT
    widths(wr, {"A": 12, "B": 11, "C": 14, "D": 16, "E": 12, "F": 18, "G": 14, "H": 16})

    # ---------------------------------------------------------------- Lift
    wl = wb.create_sheet("Lift")
    title(wl, "Lift and cumulative gains by risk decile (test set)", "Values: customers and churners per decile. Rate, lift and capture are formulas.")
    lift = pd.read_csv(REPORTS / "lift_table.csv")
    header(wl, 4, ["Decile", "Customers", "Churners", "Churn rate", "Lift", "Cumulative share of churners"])
    for i, r in enumerate(lift.itertuples(), 5):
        wl.cell(i, 1, int(r.decile))
        wl.cell(i, 2, int(r.n)).number_format = INT
        wl.cell(i, 3, int(r.churners)).number_format = INT
        wl.cell(i, 4, f"=C{i}/B{i}").number_format = PCT
        wl.cell(i, 5, f"=D{i}/(SUM($C$5:$C$14)/SUM($B$5:$B$14))").number_format = "0.00"
        wl.cell(i, 6, f"=SUM($C$5:C{i})/SUM($C$5:$C$14)").number_format = PCT
    bc = BarChart()
    bc.title, bc.height, bc.width = "Lift by risk decile", 7, 12
    bc.add_data(Reference(wl, min_col=5, min_row=4, max_row=14), titles_from_data=True)
    bc.set_categories(Reference(wl, min_col=1, min_row=5, max_row=14))
    wl.add_chart(bc, "H4")
    widths(wl, {"A": 8, "B": 11, "C": 10, "D": 11, "E": 8, "F": 16})

    # ---------------------------------------------------------------- Summary
    wsum = wb.create_sheet("Summary", 0)
    title(wsum, "Telco churn: summary", "Real data: IBM Telco Customer Churn sample. Every value is a formula over the other sheets.")
    cr = contract_rows
    rng = lambda col: "+".join(f"Segments!{col}{r}" for r in cr)
    kpis = [
        ("Customers", f"={rng('C')}", INT),
        ("Churned customers", f"={rng('D')}", INT),
        ("Churn rate", "=B6/B5", PCT2),
        ("Monthly revenue lost to churn ($)", f"={rng('E')}", USD),
        ("Annualised revenue lost ($)", "=B8*12", USD),
        ("Recommended campaign: customers targeted", "='ROI Calculator'!B13", INT),
        ("Recommended campaign: net value ($)", "='ROI Calculator'!B22", USD),
        ("Recommended campaign: ROI", "='ROI Calculator'!B23", PCT),
        ("Recommended campaign: break-even save rate", "='ROI Calculator'!B24", PCT),
        ("High-risk band: expected MRR at risk ($)", "='Risk bands'!F5", USD),
    ]
    header(wsum, 4, ["KPI", "Value"])
    for r, (lab, fml, fmt) in enumerate(kpis, 5):
        wsum.cell(r, 1, lab)
        c = wsum.cell(r, 2, fml)
        c.number_format, c.font = fmt, Font(bold=True, size=12)
    wsum.cell(16, 1, "Campaign rows reflect the ROI Calculator as saved (model group, default assumptions).").font = Font(italic=True, size=9)
    widths(wsum, {"A": 46, "B": 18})
    wsum.sheet_view.showGridLines = False

    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL_PATH)
    print(f"excel: wrote {EXCEL_PATH.name} with {len(wb.sheetnames)} sheets")


if __name__ == "__main__":
    build()
