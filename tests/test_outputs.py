"""Published outputs (dashboard, Excel, README, notebooks) agree with facts.json."""
import re
from pathlib import Path

import nbformat
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def test_dashboard_totals_match_facts(facts):
    d, k = facts["dashboard"], facts["analysis"]["kpi"]
    assert d["customers"] == k["customers"]
    assert d["churned"] == k["churned"]
    assert d["mrr_lost"] == pytest.approx(k["mrr_lost"], abs=0.5)


def test_dashboard_is_self_contained():
    html = (ROOT / "dashboard" / "index.html").read_text(encoding="utf-8")
    assert "<script src=" not in html
    assert (ROOT / "docs" / "index.html").read_bytes() == (ROOT / "dashboard" / "index.html").read_bytes()


def test_excel_summary_is_formula_driven(facts):
    ws = load_workbook(ROOT / "excel" / "retention_campaign_roi.xlsx")["Summary"]
    for row in range(5, 15):
        assert str(ws.cell(row, 2).value).startswith("=")
    if facts["excel_check"]["recalculated_by_excel"]:
        assert facts["excel_check"]["mismatches"] == 0


def test_excel_calculator_has_input_validation():
    ws = load_workbook(ROOT / "excel" / "retention_campaign_roi.xlsx")["ROI Calculator"]
    refs = " ".join(str(dv.sqref) for dv in ws.data_validations.dataValidation)
    for cell in ["B6", "B7", "B8", "B9", "B10"]:
        assert cell in refs


def test_readme_rendered_and_figures_exist():
    text = (ROOT / "README.md").read_text(encoding="utf-8")
    assert "{{" not in text and "TODO" not in text
    for path in re.findall(r"\]\((reports/figures/[^)]+)\)", text):
        assert (ROOT / path).exists(), path


@pytest.mark.parametrize("name", ["01_audit", "02_cleaning", "03_eda", "04_analysis"])
def test_notebooks_executed_without_errors(name):
    nb = nbformat.read(ROOT / "notebooks" / f"{name}.ipynb", as_version=4)
    code = [c for c in nb.cells if c.cell_type == "code"]
    assert code and all(c.execution_count for c in code)
    assert not any(o.get("output_type") == "error" for c in code for o in c.outputs)
